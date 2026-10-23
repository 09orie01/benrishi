#!/usr/bin/env python3
"""特許庁「類似商品・役務審査基準」のページに自動でアクセスできるかを確かめる試験用スクリプト。

やること
  1. 一覧ページ（索引）を取得する
  2. 最新版のページへのリンクを見つける
  3. その版のページから Excel のリンクを見つけてダウンロードする
  4. Excel の構造（シート、列、類似群コードらしい列）を表示する

結果は標準出力と out/summary.json に残します。ダウンロードした Excel も out/ に保存します。
どこかで失敗したら、終了コード 1 で終わります（GitHub Actions の画面が赤くなります）。
"""
import json
import os
import re
import sys
import time
import urllib.error
import urllib.request
from urllib.parse import urljoin

INDEX_URL = os.environ.get(
    "INDEX_URL",
    "https://www.jpo.go.jp/system/laws/rule/guideline/trademark/ruiji_kijun/index.html",
)
OUT_DIR = os.environ.get("OUT_DIR", "out")
# 自動アクセスであることを名乗ります。連絡先には、このリポジトリのURLを入れています。
USER_AGENT = os.environ.get(
    "USER_AGENT", "ruiji-kijun-check/0.1 (personal tool, monthly check; +https://github.com/09orie01/benrishi)"
)
WAIT_SECONDS = float(os.environ.get("WAIT_SECONDS", "1.5"))

VERSION_RE = re.compile(r'href="([^"]*ruiji_kijun(\d+)-(\d{4})\.html)"', re.I)
EXCEL_RE = re.compile(r'href="([^"]+\.xlsx?)"', re.I)
CODE_RE = re.compile(r"\b\d{2}[A-Z]\d{2}\b")


class CheckError(Exception):
    pass


def fetch(url):
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT, "Accept-Language": "ja"})
    try:
        with urllib.request.urlopen(req, timeout=60) as res:
            return res.status, res.read()
    except urllib.error.HTTPError as e:
        hint = "（アクセスを拒否された可能性があります）" if e.code in (403, 429) else ""
        raise CheckError(f"HTTP {e.code} {url} {hint}")
    except urllib.error.URLError as e:
        raise CheckError(f"接続できませんでした {url} ({e.reason})")


def find_versions(html):
    """索引ページから、版ごとのページへのリンクを出てくる順に集める。"""
    found = []
    seen = set()
    for m in VERSION_RE.finditer(html):
        href, ver, year = m.group(1), int(m.group(2)), int(m.group(3))
        if href in seen:
            continue
        seen.add(href)
        found.append({"href": href, "version": ver, "year": year})
    return found


def pick_latest(versions):
    """版名の数字（年、版番号）が最大のものを最新とする。ページ上の先頭と違えば知らせる。"""
    if not versions:
        raise CheckError("索引ページから版ごとのリンクを見つけられませんでした（ページの作りが変わった可能性があります）")
    latest = max(versions, key=lambda v: (v["year"], v["version"]))
    first = versions[0]
    return latest, first


def inspect_excel(path):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows = []
        code_cols = {}
        count = 0
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            count = r_idx
            if r_idx <= 6:
                rows.append([("" if c is None else str(c))[:40] for c in row])
            if r_idx <= 400:
                for c_idx, cell in enumerate(row):
                    if cell is not None and CODE_RE.search(str(cell)):
                        code_cols[c_idx + 1] = code_cols.get(c_idx + 1, 0) + 1
        sheets.append(
            {
                "name": ws.title,
                "rows": count,
                "columns": ws.max_column,
                "head": rows,
                "code_like_columns": {str(k): v for k, v in sorted(code_cols.items())},
            }
        )
    wb.close()
    return sheets


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    summary = {"index_url": INDEX_URL, "ok": False}
    try:
        print(f"[1/4] 索引ページを取得: {INDEX_URL}")
        status, body = fetch(INDEX_URL)
        html = body.decode("utf-8", errors="replace")
        print(f"      HTTP {status} / {len(html)}文字")
        versions = find_versions(html)
        latest, first = pick_latest(versions)
        summary["versions_found"] = len(versions)
        summary["latest"] = latest
        print(f"[2/4] 版のリンクを {len(versions)} 件見つけました。最新: 第{latest['version']}-{latest['year']}版 ({latest['href']})")
        if first["href"] != latest["href"]:
            print(f"      注意: ページ先頭のリンク({first['href']})と、版名から決めた最新が違います")
            summary["top_differs"] = first
        time.sleep(WAIT_SECONDS)

        version_url = urljoin(INDEX_URL, latest["href"])
        print(f"[3/4] 版のページを取得: {version_url}")
        status, body = fetch(version_url)
        vhtml = body.decode("utf-8", errors="replace")
        excel_links = [urljoin(version_url, h) for h in EXCEL_RE.findall(vhtml)]
        excel_links = list(dict.fromkeys(excel_links))
        summary["excel_links"] = excel_links
        if not excel_links:
            raise CheckError("版のページから Excel のリンクを見つけられませんでした")
        print(f"      Excel のリンク: {excel_links}")
        time.sleep(WAIT_SECONDS)

        excel_url = excel_links[0]
        print(f"[4/4] Excel をダウンロード: {excel_url}")
        status, data = fetch(excel_url)
        path = os.path.join(OUT_DIR, os.path.basename(excel_url.split("?")[0]) or "master.xlsx")
        with open(path, "wb") as f:
            f.write(data)
        summary["excel_file"] = path
        summary["excel_bytes"] = len(data)
        print(f"      保存: {path} ({len(data)} バイト)")

        sheets = inspect_excel(path)
        summary["sheets"] = sheets
        print("\n--- Excel の構造 ---")
        for s in sheets:
            print(f"シート「{s['name']}」 {s['rows']}行 × {s['columns']}列")
            print(f"  類似群コードらしい値が多い列（列番号: 件数）: {s['code_like_columns']}")
            for i, row in enumerate(s["head"], start=1):
                print(f"  {i}行目: {row}")
        summary["ok"] = True
    except CheckError as e:
        summary["error"] = str(e)
        print(f"\n失敗: {e}", file=sys.stderr)
    except Exception as e:  # 想定外の失敗も記録する
        summary["error"] = f"{type(e).__name__}: {e}"
        print(f"\n想定外の失敗: {summary['error']}", file=sys.stderr)

    with open(os.path.join(OUT_DIR, "summary.json"), "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)
    print("\n結果:", "成功" if summary["ok"] else "失敗")
    return 0 if summary["ok"] else 1


if __name__ == "__main__":
    sys.exit(main())
